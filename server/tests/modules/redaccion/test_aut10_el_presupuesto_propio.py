"""AUT.10 (issue #120) — el presupuesto propio como dos funciones de tarea, de punta a punta.

Es la prueba de aceptación del hito «funciones de tarea»: un cuaderno real convertido en
funciones del catálogo, que usa lo que el hito construyó —varios ficheros de entrada (#194),
artefactos de salida (AUT.7)— y lo que ya había: el contrato, la auditoría y el registro.

Las dos funciones viven en `docs/ejemplos/funciones/presupuesto/`, con sus contratos y un
generador de datos sintéticos. **Nunca el CSV real**: lleva nombres y NIF de terceros.

1. `anonimizar_presupuesto` — la anonimización del cuaderno, registrada como función aparte.
2. `documento_presupuesto` — normaliza lo ya anonimizado, agrega y produce el CSV limpio y el Word.

Esto corre por el **modo local** (pandas 2). El sandbox de verdad corre con pandas 3, y por eso
hay una prueba gemela en `services/script_sandbox/tests/test_aut10_el_presupuesto_propio.py`:
un guion que funciona con una versión y no con la otra pasaría aquí y fallaría desplegado.
"""
from __future__ import annotations

import importlib.util
import io
import json
import sys
import uuid
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parents[4]
EJEMPLO = RAIZ / "docs" / "ejemplos" / "funciones" / "presupuesto"


def _modulo(ruta: Path, nombre: str):
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    assert spec and spec.loader
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[nombre] = modulo
    spec.loader.exec_module(modulo)
    return modulo


DATOS = _modulo(EJEMPLO / "datos_sinteticos.py", "_aut10_datos_sinteticos")
CONTRATOS = json.loads((EJEMPLO / "contratos.json").read_text(encoding="utf-8"))


def _guion(clave: str) -> str:
    return (EJEMPLO / CONTRATOS[clave]["guion"]).read_text(encoding="utf-8")


class _Almacen:
    """Un `StorageService` en memoria."""

    def __init__(self, **documentos: bytes) -> None:
        self.contenido = dict(documentos)

    async def put(self, key: str, data: bytes) -> None:
        self.contenido[key] = data

    async def get(self, key: str) -> bytes:
        if key not in self.contenido:
            raise FileNotFoundError(key)
        return self.contenido[key]

    async def delete(self, key: str) -> None:
        self.contenido.pop(key, None)

    async def exists(self, key: str) -> bool:
        return key in self.contenido


def _pat(organizacion):
    from server.app.core.auth.models import UserInfo

    return UserInfo(
        user_id=f"pat-{uuid.uuid4()}",
        email="pressupost@example.test",
        role="admin",
        organizacion_ids=[str(organizacion)],
    )


async def _registrar(db_session, clave: str, organizacion):
    from server.app.modules.redaccion.contracts.funciones import ContratoFuncion
    from server.app.modules.redaccion.funciones_service import registrar_version

    datos = CONTRATOS[clave]
    funcion, version = await registrar_version(
        db_session,
        nombre=f"{datos['nombre']} {uuid.uuid4().hex[:6]}",
        organizacion_id=organizacion,
        code=_guion(clave),
        contrato=ContratoFuncion.model_validate(datos["contrato"]),
        declarada_por=uuid.uuid4(),
    )
    version.estado = "registrada"
    db_session.add(version)
    await db_session.flush()
    return funcion


async def _ejecutar(db_session, funcion, organizacion, almacen, ficheros, parametros=None):
    from server.app.core.sandbox_client import LocalSandboxClient
    from server.app.routers.redaccion.funciones_run_router import (
        EjecutarRequest,
        ejecutar_funcion_por_api,
    )

    return await ejecutar_funcion_por_api(
        funcion_id=funcion.id,
        body=EjecutarRequest(version=1, ficheros=ficheros, parametros=parametros or {}),
        principal=_pat(organizacion),
        session=db_session,
        sandbox=LocalSandboxClient(),
        almacen=almacen,
    )


def _artefacto(almacen: _Almacen, respuesta, nombre: str) -> bytes:
    """El contenido de un artefacto de la respuesta, leído de donde la plataforma lo guardó."""
    artefacto = next(a for a in respuesta.artefactos if a.nombre == nombre)
    clave = next(k for k in almacen.contenido if k.endswith(f"/{artefacto.nombre}"))
    return almacen.contenido[clave]


class TestLasFuncionesSonValidas:
    """Antes de ejecutar nada: lo que se registraría tiene que pasar lo que la plataforma exige."""

    @pytest.mark.parametrize("clave", list(CONTRATOS))
    def test_el_contrato_se_sostiene(self, clave):
        from server.app.modules.redaccion.contracts.funciones import ContratoFuncion

        ContratoFuncion.model_validate(CONTRATOS[clave]["contrato"])

    @pytest.mark.parametrize("clave", list(CONTRATOS))
    def test_el_auditor_de_la_api_no_encuentra_nada(self, clave):
        """Ni críticos ni avisos: el sandbox rechaza también los avisos."""
        from server.app.modules.redaccion.services.script_auditor import ScriptSecurityAuditor

        resultado = ScriptSecurityAuditor().audit(_guion(clave))

        assert resultado.findings == [], [f.message for f in resultado.findings]

    @pytest.mark.parametrize("clave", list(CONTRATOS))
    def test_y_el_del_sandbox_tampoco(self, clave):
        """La segunda barrera, la que corre justo antes del `exec`, es más estricta a propósito."""
        auditor = _modulo(
            RAIZ / "services" / "script_sandbox" / "sandbox" / "auditor.py",
            "_aut10_auditor_del_sandbox",
        )

        resultado = auditor.ScriptSecurityAuditor().audit(_guion(clave))

        assert resultado.approved, resultado.findings


class TestAnonimizar:

    @pytest.mark.asyncio
    async def test_no_sobrevive_ninguna_persona(self, db_session):
        import pandas as pd

        organizacion = uuid.uuid4()
        funcion = await _registrar(db_session, "anonimizar_presupuesto", organizacion)
        almacen = _Almacen(**{"org/p.csv": DATOS.presupuesto(), "org/pdi.txt": DATOS.pdi()})

        respuesta = await _ejecutar(
            db_session, funcion, organizacion, almacen,
            {"presupuesto": "org/p.csv", "pdi": "org/pdi.txt"},
        )

        salida = _artefacto(almacen, respuesta, "presupuesto_anonimizado.csv").decode("utf-8-sig")
        for persona in DATOS.PERSONAS:
            assert persona not in salida, f"«{persona}» sigue en el CSV anonimizado"
        assert "[IP - dada protegida]" in salida

        # Y lo que no es una persona sigue como estaba: tachar de más también es un fallo.
        for intacto in DATOS.NO_PERSONAS:
            assert intacto in salida, f"«{intacto}» no es una persona y ha desaparecido"

        tabla = pd.read_csv(io.StringIO(salida), dtype=str, keep_default_na=False)
        gastos = tabla[tabla["Pcr Apl Tipo"] == "P"]
        ingresos = tabla[tabla["Pcr Apl Tipo"] == "C"]
        assert (gastos["Tercero Nif"] == "").all(), "un NIF de gasto ha sobrevivido"
        assert DATOS.NIF_ENTIDAD in set(ingresos["Tercero Nif"]), "el NIF de la entidad se ha perdido"

    @pytest.mark.asyncio
    async def test_dos_personas_en_la_misma_descripcion(self, db_session):
        """El cuaderno sustituía una vez; la segunda persona quedaba a la vista."""
        import pandas as pd

        organizacion = uuid.uuid4()
        funcion = await _registrar(db_session, "anonimizar_presupuesto", organizacion)
        almacen = _Almacen(**{"org/p.csv": DATOS.presupuesto(), "org/pdi.txt": DATOS.pdi()})

        respuesta = await _ejecutar(
            db_session, funcion, organizacion, almacen,
            {"presupuesto": "org/p.csv", "pdi": "org/pdi.txt"},
        )

        salida = _artefacto(almacen, respuesta, "presupuesto_anonimizado.csv").decode("utf-8-sig")
        tabla = pd.read_csv(io.StringIO(salida), dtype=str, keep_default_na=False)
        doble = tabla[tabla["Nombre Subproyecto"].str.startswith("PID2022-1234")].iloc[0]
        assert doble["Nombre Subproyecto"].count("[IP - dada protegida]") == 2


class TestElDocumento:

    @pytest.mark.asyncio
    async def test_de_punta_a_punta_sobre_lo_anonimizado(self, db_session):
        """Las dos funciones encadenadas, que es como se usan: la salida de una es la entrada de la otra."""
        import pandas as pd
        from docx import Document

        organizacion = uuid.uuid4()
        anonimizar = await _registrar(db_session, "anonimizar_presupuesto", organizacion)
        documento = await _registrar(db_session, "documento_presupuesto", organizacion)
        almacen = _Almacen(**{
            "org/p.csv": DATOS.presupuesto(),
            "org/pdi.txt": DATOS.pdi(),
            "org/clasif.csv": DATOS.clasificacion(),
            "org/prosa.md": DATOS.prosa(),
        })

        primera = await _ejecutar(
            db_session, anonimizar, organizacion, almacen,
            {"presupuesto": "org/p.csv", "pdi": "org/pdi.txt"},
        )
        almacen.contenido["org/anon.csv"] = _artefacto(
            almacen, primera, "presupuesto_anonimizado.csv"
        )

        segunda = await _ejecutar(
            db_session, documento, organizacion, almacen,
            {
                "presupuesto": "org/anon.csv",
                "clasificacion": "org/clasif.csv",
                "prosa": "org/prosa.md",
            },
            parametros={"titulo": "Pressupost sintètic 2026"},
        )

        assert {a.nombre for a in segunda.artefactos} == {"pressupost_net.csv", "pressupost_uji.docx"}
        metricas = {m["name"]: m["value"] for m in segunda.metrics}
        assert metricas["total_gastos"] == pytest.approx(DATOS.total_gastos_del_ultimo_ejercicio())
        assert metricas["secciones_con_prosa"] == 2

        # El CSV limpio cuadra con lo que dice el documento.
        neto = pd.read_csv(
            io.StringIO(_artefacto(almacen, segunda, "pressupost_net.csv").decode("utf-8-sig")),
            dtype=str, keep_default_na=False,
        )
        assert "cap_desc" in neto.columns
        for persona in DATOS.PERSONAS:
            assert not neto.apply(lambda c: c.str.contains(persona, regex=False)).any().any()

        # El Word se abre, lleva el título pedido y la prosa con las cifras puestas.
        word = Document(io.BytesIO(_artefacto(almacen, segunda, "pressupost_uji.docx")))
        texto = "\n".join(p.text for p in word.paragraphs)
        assert "Pressupost sintètic 2026" in texto
        assert "El pressupost de despeses puja a" in texto
        assert "{{VAL:" not in texto, "un marcador de cifra se ha quedado sin sustituir"
        assert "‹?" not in texto, "la prosa pide una cifra que el documento no calcula"
        assert len(word.tables) >= 5


def _funciones_del_documento() -> dict:
    """Las funciones de `documento_presupuesto`, sin la parte que lee `options` y ejecuta."""
    codigo = _guion("documento_presupuesto")
    definiciones = codigo.split('\nrutas = options["ficheros"]', 1)[0]
    espacio: dict = {}
    exec(compile(definiciones, "documento_presupuesto.py", "exec"), espacio)
    return espacio


def test_a_subset_that_adds_up_to_zero_reports_zero_not_one_euro():
    # Copilot en la PR #212: `total = suma or 1.0` servía de divisor y de total a la vez, así que
    # un subconjunto sin gasto salía en el Word con una fila TOTAL de 1,00 €.
    import pandas as pd

    taula_concepte = _funciones_del_documento()["taula_concepte"]
    vacio = pd.DataFrame(
        {
            "cap": ["2"], "cap_desc": ["Gastos corrientes"], "art": ["22"],
            "art_desc": ["Material"], "conc": ["220"], "conc_desc": ["Oficina"],
            "importe": [0.0],
        }
    )

    tabla = taula_concepte(vacio)

    total = tabla[tabla["tipus_fila"] == "total"].iloc[0]
    assert total["import"] == 0
