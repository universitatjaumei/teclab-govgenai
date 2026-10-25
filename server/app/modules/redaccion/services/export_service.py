"""ExportService — exportación DOCX del informe ensamblado (9R.9.2 / 1C.4).

Deploy: edge
Alcance MVP: DOCX únicamente. ODT está en backlog post-MVP para no romper
compatibilidad de estilos entre LibreOffice y Word; la abstracción `Exporter`
se conserva para que entre como segunda implementación sin rework.

Dependencia: requiere DraftingRunManifest con final_document_hash no nulo.
Sin hash el documento no está ensamblado y la exportación se rechaza.
"""
from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.shared import Inches, RGBColor

from server.app.modules.redaccion.contracts.manifest import DraftingRunManifest


class ExportNotReadyError(Exception):
    """Raised when manifest.final_document_hash is None."""


class ExportService:
    """Genera un DOCX a partir del contenido ensamblado y el manifest de auditoría."""

    def __init__(self, builder=None) -> None:
        self._builder = builder
        self.last_payload_used = None

    async def export(self, workspace_id, format: str = "docx") -> bytes:  # noqa: ARG002 — DOCX es el único formato hoy
        """Exporta usando PreviewBuilderService (1C.3).

        Requiere que se haya inyectado un builder en el constructor.
        """
        if self._builder is None:
            raise RuntimeError("ExportService requires a PreviewBuilderService builder for export()")
        payload = await self._builder.build_payload(workspace_id)
        self.last_payload_used = payload
        return self._payload_to_docx(payload)

    def _payload_to_docx(self, payload) -> bytes:
        """Genera DOCX a partir de un PreviewPayload.

        PRO.5 — esto hacía `doc.add_paragraph(block.html)`, o sea que metía el HTML **como
        texto**: desde PRO.3, cuando la vista previa empezó a pintar las tablas extraídas, el
        documento exportado llevaba dentro `<table><thead><tr><th>capitulo…` escrito a mano.
        Con el contenido estructurado del bloque, una tabla se escribe como tabla y un gráfico
        se incrusta como imagen; el HTML queda para el navegador.
        """
        doc = Document()
        doc.add_heading(payload.cover.title, level=0)
        for section in payload.body:
            doc.add_heading(section.title, level=1)
            for block in section.blocks:
                self._escribir_bloque(doc, block)
        if payload.audit_annex:
            doc.add_page_break()
            doc.add_heading("Anexo de Auditoría", level=1)
            tbl = doc.add_table(rows=1, cols=3)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Chunk"
            hdr[1].text = "Fuente"
            hdr[2].text = "Página"
            for entry in payload.audit_annex:
                row = tbl.add_row().cells
                row[0].text = str(entry.chunk_id)
                row[1].text = entry.source_filename or entry.source_url or "—"
                row[2].text = str(entry.page) if entry.page is not None else "—"
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def export_to_docx(
        self,
        manifest: DraftingRunManifest,
        document_content: str,
    ) -> bytes:
        """Exporta el documento ensamblado a DOCX.

        Args:
            manifest: DraftingRunManifest con final_document_hash no nulo.
            document_content: Texto del documento ensamblado (markdown/plain text).

        Returns:
            Bytes del DOCX generado.

        Raises:
            ExportNotReadyError: si manifest.final_document_hash is None.
        """
        if manifest.final_document_hash is None:
            raise ExportNotReadyError(
                "Cannot export: manifest.final_document_hash is None. "
                "Run the DraftingCoreGraph to completion before exporting."
            )

        doc = Document()
        self._write_content(doc, document_content)
        self._write_audit_appendix(doc, manifest)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    def _escribir_bloque(self, doc: DocumentoDocx, block) -> None:
        """Un bloque del informe en el documento, según lo que sea.

        Sin contenido estructurado se cae al HTML como texto, que es el comportamiento
        anterior: un bloque viejo no puede impedir la exportación.
        """
        contenido = getattr(block, "content", None)
        if not isinstance(contenido, dict):
            if block.html:
                doc.add_paragraph(block.html)
            return

        texto = contenido.get("text") or contenido.get("value")
        if texto:
            doc.add_paragraph(str(texto))

        grafico = contenido.get("chart")
        if isinstance(grafico, dict):
            imagen = (getattr(block, "images", None) or {}).get(grafico.get("storage_key", ""))
            if imagen:
                doc.add_picture(io.BytesIO(imagen), width=Inches(6))

        for tabla in contenido.get("tables") or []:
            self._escribir_tabla(doc, tabla)

        for metrica in contenido.get("metrics") or []:
            unidad = f" {metrica['unit']}" if metrica.get("unit") else ""
            doc.add_paragraph(
                f"{metrica.get('name', '')}: {metrica.get('value', '')}{unidad}",
                style="List Bullet",
            )

        if contenido.get("free_text"):
            doc.add_paragraph(str(contenido["free_text"]))

    @staticmethod
    def _escribir_tabla(doc: DocumentoDocx, tabla: dict) -> None:
        cabeceras = [str(c) for c in (tabla.get("headers") or [])]
        filas = tabla.get("rows") or []
        if not cabeceras and not filas:
            return
        if tabla.get("name"):
            doc.add_paragraph(str(tabla["name"]))
        columnas = len(cabeceras) or max((len(f) for f in filas), default=1)
        tabla_docx = doc.add_table(rows=1 if cabeceras else 0, cols=columnas)
        tabla_docx.style = "Table Grid"
        if cabeceras:
            for i, cabecera in enumerate(cabeceras):
                tabla_docx.rows[0].cells[i].text = cabecera
        for fila in filas:
            celdas = tabla_docx.add_row().cells
            for i, valor in enumerate(fila[:columnas]):
                celdas[i].text = str(valor)

    def _write_content(self, doc: DocumentoDocx, content: str) -> None:
        doc.add_heading("Informe Generado", level=0)
        for line in content.splitlines():
            doc.add_paragraph(line)

    def _write_audit_appendix(self, doc: DocumentoDocx, manifest: DraftingRunManifest) -> None:
        doc.add_page_break()
        heading = doc.add_heading("Anexo de Auditoría", level=1)
        heading.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        doc.add_paragraph(
            f"Hash del documento: {manifest.final_document_hash}",
        )
        doc.add_paragraph(
            f"Perfil: {manifest.report_profile}  |  "
            f"Estado al cierre: {manifest.status_at_close}  |  "
            f"Generado: {manifest.created_at.strftime('%Y-%m-%d %H:%M UTC') if manifest.created_at else '—'}",
        )

        if manifest.uploaded_documents:
            doc.add_heading("Documentos subidos", level=2)
            tbl = doc.add_table(rows=1, cols=3)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Slot"
            hdr[1].text = "Archivo"
            hdr[2].text = "Tamaño (bytes)"
            for d in manifest.uploaded_documents:
                row = tbl.add_row().cells
                row[0].text = d.slot_id
                row[1].text = d.filename
                row[2].text = str(d.size_bytes)

        if manifest.ai_blocks:
            doc.add_heading("Bloques IA", level=2)
            tbl = doc.add_table(rows=1, cols=4)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Bloque"
            hdr[1].text = "Tipo"
            hdr[2].text = "Modelo"
            hdr[3].text = "Versión prompt"
            for b in manifest.ai_blocks:
                row = tbl.add_row().cells
                row[0].text = b.block_id
                row[1].text = b.kind
                row[2].text = b.model_used or "—"
                row[3].text = b.prompt_version or "—"

        if manifest.user_approvals:
            doc.add_heading("Aprobaciones", level=2)
            tbl = doc.add_table(rows=1, cols=3)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Aprobado por"
            hdr[1].text = "Fecha"
            hdr[2].text = "Nota"
            for a in manifest.user_approvals:
                row = tbl.add_row().cells
                row[0].text = str(a.approved_by)
                row[1].text = a.approved_at.strftime("%Y-%m-%d %H:%M UTC")
                row[2].text = a.note or "—"

        if manifest.warnings:
            doc.add_heading("Advertencias de extracción", level=2)
            tbl = doc.add_table(rows=1, cols=3)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Bloque"
            hdr[1].text = "Tipo"
            hdr[2].text = "Mensaje"
            for w in manifest.warnings:
                row = tbl.add_row().cells
                row[0].text = w.block_id or "—"
                row[1].text = w.kind
                row[2].text = w.message
