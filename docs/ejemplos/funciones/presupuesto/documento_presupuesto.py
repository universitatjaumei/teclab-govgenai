# El presupuesto propio: CSV limpio y Word maquetado (AUT.10, issue #120).
#
# Segunda de las dos funciones en que se convierte el cuaderno «Pressupost UJI». Recibe el
# presupuesto **ya anonimizado** por `anonimizar_presupuesto.py` y hace lo que el cuaderno hacía
# después:
#
#   Fase 1  normaliza (el port a pandas del SQL del dataset de Superset) y cruza con la
#           clasificación económica;
#   Fase 3  agrega las tablas del documento;
#   Fase 4  escribe el CSV limpio (`pressupost_net.csv`) y el Word (`pressupost_uji.docx`).
#
# La Fase 2 —anonimizar— **no está aquí**: es la otra función, y se ejecuta antes. Separarla es lo
# que permite usarla sola y saber, mirando el registro, que este documento se hizo sobre datos ya
# anonimizados.
#
# **La prosa ya no se lee de un Google Doc.** El cuaderno montaba la unidad de la persona y leía
# el documento con su cuenta, que es lo que la regla de soberanía local prohíbe centralizar. Ahora
# es un Markdown que se sube, con los mismos marcadores: el texto de cada sección va antes de su
# `{{TAULA:<id>:INICI}}`, y `{{VAL:<clave>}}` se sustituye por la cifra vigente.
#
# Cambios de forma respecto al cuaderno, ninguno de resultado: no hay `pip install` (python-docx
# viene en la imagen), los ficheros se leen con pandas porque el auditor deniega `open`, y el tipo
# numérico se comprueba con numpy en vez de con `numbers`, que no está en el ecosistema.
import re

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

# =============================================================================
#  FASE 1 — NORMALIZACIÓN
# =============================================================================
RENAMES = {
    "Presupuesto": "ejercicio", "Pcr Apl Tipo": "apl_tipo",
    "Solicitud Prespuesto": "solicitud", "Af": "af",
    "Centre": "centro_cod", "Descr Centre": "centro_desc",
    "Sbc": "subcentro_cod", "Descr Sbc": "subcentro_desc",
    "Tipo Linea": "tipo_linea", "Pcr Lfi Id": "lfi_id", "Descr Lfi": "lfi_desc",
    "Pcr Apl Id": "aplicacion_id", "Tipo Proyecto": "tipo_proyecto",
    "Proj": "proyecto_cod", "Descr Proj": "proyecto_desc",
    "Sbp": "subproyecto_cod", "Nombre Subproyecto": "subproyecto_desc",
    "Pro G I": "pgi", "Pcr Spy Id": "spy_id", "Pcr Prg Id": "programa_funcional",
    "Cap": "cap", "Art": "art", "Conc": "conc", "Pcr Importe SUM": "importe",
    "Observacion": "observacion", "Comentario Centro": "comentario_centro",
    "Etiquetas Ods": "etiquetas_ods_raw", "Etiquetas Generos": "etiquetas_generos",
    "Etiquetas Generales": "etiquetas_generales",
    "Expediente Pluri Nombre": "expediente_pluri_nombre",
    "Expediente Pluri Descripcion": "expediente_pluri_desc",
    "Pcr Sub Per Id": "sub_per_id", "Tercero Nombre": "tercero_nombre",
    "Tercero Nif": "tercero_nif", "Tercero Per Id": "tercero_per_id",
    "Subpro Id Origen Financiacion": "subpro_origen_financ",
    "Pro Id Origen Financiacion": "pro_origen_financ",
    "Expedinete Cma": "expediente_cma", "Pcr Id": "pcr_id",
}

PROGRAMA_MAP = {
    "422-A": "422-E - Universitat i estudis superiors",
    "422-C": "422-E - Universitat i estudis superiors",
    "422-D": "422-E - Universitat i estudis superiors",
    "541-A": "542-C - Investigació, desenvolupament tecnològic i innovació (I+D+i)",
}

THRESHOLD_TOP = 500000
PREPOSICIONS = ["De", "Del", "Dels", "Des", "La", "Las", "Les", "El", "Els",
                "Los", "Y", "En", "Amb", "Con", "Per", "Para", "Por"]
SIGLES = {"Sepie": "SEPIE", "Cf": "CF", "Cv": "CV"}


def fase1_normalitza(raw_csv, clasif_csv):
    df = pd.read_csv(raw_csv, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df = df.rename(columns=RENAMES)
    df = df[[c for c in RENAMES.values() if c in df.columns]].copy()

    df["ejercicio"] = pd.to_numeric(df["ejercicio"], errors="coerce").astype("Int64")
    # Redondeo a céntimos en origen: el presupuesto es exacto al céntimo, así que esto sólo quita
    # el ruido de coma flotante.
    df["importe"] = pd.to_numeric(
        df["importe"].str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
        if df["importe"].str.contains(r"\.\d{3},", regex=True).any()
        else df["importe"].str.replace(",", ".", regex=False),
        errors="coerce").fillna(0.0).round(2)
    df["apl_tipo_desc"] = df["apl_tipo"].map({"P": "Despeses", "C": "Ingressos"}).fillna(df["apl_tipo"])
    df["tipo_clasif"] = df["apl_tipo"].map({"P": "D", "C": "I"}).fillna(df["apl_tipo"])
    df["es_despesa"] = df["apl_tipo"].eq("P")
    df["es_ingres"] = df["apl_tipo"].eq("C")
    df["af_desc"] = df["af"].map({"S": "Afectat", "N": "No afectat"}).fillna(df["af"])
    df["programa_desc"] = df["programa_funcional"].map(PROGRAMA_MAP).fillna(df["programa_funcional"])

    maxej = df["ejercicio"].max()
    df["es_actual"] = df["ejercicio"].eq(maxej)
    df["es_anterior"] = df["ejercicio"].eq(maxej - 1)

    dic = pd.read_csv(clasif_csv, dtype=str, keep_default_na=False)
    look = {(r.tipo, r.nivel, str(r.codigo).strip()): r.descripcion for r in dic.itertuples()}

    def desc_econ(row, nivel, col):
        d = look.get((row["tipo_clasif"], nivel, str(row[col]).strip()))
        return f"{row[col]} - {d}" if d else str(row[col])

    df["cap_desc"] = df.apply(lambda r: desc_econ(r, "cap", "cap"), axis=1)
    df["art_desc"] = df.apply(lambda r: desc_econ(r, "art", "art"), axis=1)
    df["conc_desc"] = df.apply(lambda r: desc_econ(r, "conc", "conc"), axis=1)

    cen_tot = (df[df["es_despesa"] & df["es_actual"]].groupby("centro_desc")["importe"].sum())
    ter_tot = (df[df["es_ingres"] & df["es_actual"]].groupby("tercero_per_id")["importe"].sum())
    df["centro_top"] = df.apply(
        lambda r: r["centro_desc"] if (r["es_ingres"] or cen_tot.get(r["centro_desc"], 0) >= THRESHOLD_TOP)
        else "Altres centres (< 0,5 M€)", axis=1)
    return df, ter_tot


def _title_case(s):
    return re.sub(r"[A-Za-zÀ-ÿ]+", lambda m: m.group(0).capitalize(), s)


def _normalitza_entitat(nom):
    nom = re.sub(r"^(.+), \1$", r"\1", nom)
    nom = _title_case(nom)
    for p in PREPOSICIONS:
        nom = re.sub(rf" {p} ", f" {p.lower()} ", nom)
    for k, v in SIGLES.items():
        nom = re.sub(rf"\b{k}\b", v, nom)
    return nom


def terceros(df, ter_tot):
    """Lo que el cuaderno hacía en la Fase 2 **sin** anonimizar: eso ya viene hecho.

    En los ingresos el tercero es la entidad que financia y se normaliza su nombre; en los gastos
    se toma tal cual, porque la función de anonimizar ya lo tachó si era una persona.
    """
    df["proyecto_desc_anon"] = df["proyecto_desc"]
    df["subproyecto_desc_anon"] = df["subproyecto_desc"]
    df["expediente_pluri_nombre_anon"] = df["expediente_pluri_nombre"]

    def tercero_disp(r):
        if r["es_ingres"]:
            return _normalitza_entitat(r["tercero_nombre"])
        return r["tercero_nombre"]

    def tercero_top(r):
        if r["es_ingres"]:
            if ter_tot.get(r["tercero_per_id"], 0) >= THRESHOLD_TOP:
                return _normalitza_entitat(r["tercero_nombre"])
            return "Altres entitats (< 0,5 M€)"
        return r["tercero_nombre"]

    df["tercero_display"] = df.apply(tercero_disp, axis=1)
    df["tercero_top"] = df.apply(tercero_top, axis=1)
    return df


# =============================================================================
#  FASE 3 — AGREGACIÓN
# =============================================================================
PARAMS = dict(
    LLINDAR_CENTRE=500000,
    LLINDAR_TERCER=500000,
    INCLOU_A2=True,
    SUBTOTALS_ARTICLE=True,
    INCLOU_INTERANUAL=True,
)


def _caps_ordenats(d):
    caps = d[["cap", "cap_desc"]].drop_duplicates()
    caps["k"] = pd.to_numeric(caps["cap"], errors="coerce")
    return caps.sort_values("k")


def taula_concepte(d, subtotals_article=True):
    g = (d.groupby(["cap", "cap_desc", "art", "art_desc", "conc", "conc_desc"],
                   as_index=False)["importe"].sum())
    g["k_cap"] = pd.to_numeric(g["cap"], errors="coerce")
    g["k_art"] = pd.to_numeric(g["art"], errors="coerce")
    g["k_conc"] = pd.to_numeric(g["conc"], errors="coerce")
    total = g["importe"].sum() or 1.0
    rows = []
    for cap, capd in _caps_ordenats(d)[["cap", "cap_desc"]].itertuples(index=False):
        gc = g[g["cap"] == cap].sort_values(["k_art", "k_conc"])
        cap_sum = gc["importe"].sum()
        for art in gc["art"].drop_duplicates():
            ga = gc[gc["art"] == art]
            for r in ga.itertuples(index=False):
                rows.append(["detall", r.cap_desc, r.art_desc, r.conc_desc, r.importe,
                             r.importe / cap_sum if cap_sum else 0, r.importe / total])
            if subtotals_article and len(ga) > 1:
                asum = ga["importe"].sum()
                rows.append(["subtotal_art", capd, ga.iloc[0]["art_desc"], "Subtotal article",
                             asum, asum / cap_sum if cap_sum else 0, asum / total])
        rows.append(["subtotal_cap", capd, "", "Subtotal capítol", cap_sum, 1.0, cap_sum / total])
    rows.append(["total", "TOTAL", "", "", total, 1.0, 1.0])
    return pd.DataFrame(rows, columns=["tipus_fila", "capitol", "article", "concepte",
                                       "import", "pct_cap", "pct_total"])


def matriu_x_capitol(d, fila_col, ordre=None):
    piv = d.pivot_table(index=fila_col, columns="cap_desc", values="importe",
                        aggfunc="sum", fill_value=0)
    cols = _caps_ordenats(d)["cap_desc"].drop_duplicates().tolist()
    piv = piv.reindex(columns=[c for c in cols if c in piv.columns], fill_value=0)
    piv["Total"] = piv.sum(axis=1)
    if ordre is not None:
        piv = piv.reindex([o for o in ordre if o in piv.index])
    else:
        piv = piv.sort_values("Total", ascending=False)
    gtot = piv["Total"].sum() or 1.0
    piv["% Total"] = piv["Total"] / gtot
    tot = piv.drop(columns="% Total").sum(axis=0)
    tot["% Total"] = 1.0
    piv = pd.concat([piv, pd.DataFrame([tot], index=["TOTAL"])])
    return piv.reset_index().rename(columns={fila_col: "fila", "index": "fila"})


def _resum_programa(dg):
    # Sin `groupby().apply()`: su firma cambió entre pandas 2 y 3 (`include_groups`), y el guion
    # corre con las dos —2 en el modo local de la plataforma, 3 en el sandbox—. Tres sumas por
    # grupo dicen lo mismo y no dependen de la versión.
    total = dg.groupby("programa_desc")["importe"].sum()
    sense = dg[dg.af.eq("N")].groupby("programa_desc")["importe"].sum()
    amb = dg[dg.af.eq("S")].groupby("programa_desc")["importe"].sum()
    out = pd.DataFrame({"sense_afectacio": sense, "amb_afectacio": amb, "total": total}).fillna(0.0)
    out = out.reset_index().rename(columns={"index": "programa_desc"})
    out["pct_total"] = out["total"] / (out["total"].sum() or 1)
    return out


def genera_taules(net_csv, params=PARAMS):
    df = pd.read_csv(net_csv, keep_default_na=False, dtype=str, encoding="utf-8-sig")
    df["importe"] = pd.to_numeric(df["importe"], errors="coerce").fillna(0.0)
    for b in ("es_despesa", "es_ingres", "es_actual", "es_anterior"):
        df[b] = df[b].astype(str).str.lower().eq("true")
    act = df[df["es_actual"]]
    T = {}

    def tot(mask):
        return df[mask]["importe"].sum()

    des_a, des_p = tot(df.es_actual & df.es_despesa), tot(df.es_anterior & df.es_despesa)
    ing_a, ing_p = tot(df.es_actual & df.es_ingres), tot(df.es_anterior & df.es_ingres)
    T["00_Resum"] = pd.DataFrame([
        ["Total despeses", des_p, des_a, des_a - des_p, (des_a - des_p) / des_p if des_p else 0],
        ["Total ingressos", ing_p, ing_a, ing_a - ing_p, (ing_a - ing_p) / ing_p if ing_p else 0],
        ["Despeses sense afectació", "", tot(df.es_actual & df.es_despesa & df.af.eq("N")), "", ""],
        ["Despeses amb afectació", "", tot(df.es_actual & df.es_despesa & df.af.eq("S")), "", ""],
    ], columns=["concepte", "anterior", "actual", "variacio_eur", "variacio_pct"])

    dg = act[act.es_despesa]
    T["A1_1_sense_afectacio"] = taula_concepte(dg[dg.af.eq("N")], params["SUBTOTALS_ARTICLE"])
    T["A1_2_amb_afectacio"] = taula_concepte(dg[dg.af.eq("S")], params["SUBTOTALS_ARTICLE"])
    T["A1_3_total"] = taula_concepte(dg, params["SUBTOTALS_ARTICLE"])

    T["A2_0_resum_programa"] = _resum_programa(dg)
    if params["INCLOU_A2"]:
        for i, prog in enumerate(sorted(dg["programa_desc"].unique()), 1):
            dp = dg[dg.programa_desc == prog]
            T[f"A2_{i}_1_{_slug(prog)}_sense"] = taula_concepte(dp[dp.af.eq("N")], params["SUBTOTALS_ARTICLE"])
            T[f"A2_{i}_2_{_slug(prog)}_amb"] = taula_concepte(dp[dp.af.eq("S")], params["SUBTOTALS_ARTICLE"])
            T[f"A2_{i}_3_{_slug(prog)}_total"] = taula_concepte(dp, params["SUBTOTALS_ARTICLE"])

    T["A3_0_resum_centre_capitol"] = matriu_x_capitol(dg, "centro_top")
    cen_tot = dg.groupby("centro_desc")["importe"].sum()
    principals = cen_tot[cen_tot >= params["LLINDAR_CENTRE"]].sort_values(ascending=False).index.tolist()
    for i, cen in enumerate(principals, 1):
        dc = dg[dg.centro_desc == cen]
        T[f"A3_{i}_{_slug(cen)}"] = matriu_x_capitol(dc, "subcentro_desc")

    ig = act[act.es_ingres]
    T["B1_1_sense_afectacio"] = taula_concepte(ig[ig.af.eq("N")], params["SUBTOTALS_ARTICLE"])
    T["B1_2_amb_afectacio"] = taula_concepte(ig[ig.af.eq("S")], params["SUBTOTALS_ARTICLE"])
    T["B1_3_total"] = taula_concepte(ig, params["SUBTOTALS_ARTICLE"])

    T["B2_0_resum_tercer_capitol"] = matriu_x_capitol(ig, "tercero_top")
    ter_tot = ig.groupby("tercero_display")["importe"].sum()
    ter_princ = ter_tot[ter_tot >= params["LLINDAR_TERCER"]].sort_values(ascending=False).index.tolist()
    ter_princ = [t for t in ter_princ if not str(t).startswith("Altres")]
    for i, ter in enumerate(ter_princ, 1):
        dt = ig[ig.tercero_display == ter]
        T[f"B2_{i}_{_slug(ter)}"] = taula_concepte(dt, subtotals_article=False)

    if params["INCLOU_INTERANUAL"]:
        T["C1_1_var_despeses_capitol"] = _var_x(df, df.es_despesa, "cap_desc")
        T["C1_2_var_despeses_programa"] = _var_x(df, df.es_despesa, "programa_desc")
        T["C2_1_var_ingressos_capitol"] = _var_x(df, df.es_ingres, "cap_desc")
        T["C2_2_var_ingressos_tercer"] = _var_x(df, df.es_ingres, "tercero_top")

    return T


def _var_x(df, mask, col):
    d = df[mask]
    a = d[d.es_actual].groupby(col)["importe"].sum()
    p = d[d.es_anterior].groupby(col)["importe"].sum()
    out = pd.DataFrame({"anterior": p, "actual": a}).fillna(0.0)
    out["variacio_eur"] = out["actual"] - out["anterior"]
    out["variacio_pct"] = out.apply(
        lambda r: r["variacio_eur"] / r["anterior"] if r["anterior"] else 0, axis=1)
    out = out.sort_values("actual", ascending=False).reset_index().rename(columns={col: "fila", "index": "fila"})
    tot = out[["anterior", "actual", "variacio_eur"]].sum()
    tot["fila"] = "TOTAL"
    tot["variacio_pct"] = tot["variacio_eur"] / tot["anterior"] if tot["anterior"] else 0
    return pd.concat([out, pd.DataFrame([tot])], ignore_index=True)


def _slug(s):
    s = re.sub(r"[^A-Za-z0-9]+", "_", str(s)).strip("_")
    return s[:22]


# =============================================================================
#  FASE 4 — EL WORD
# =============================================================================
SECTIONS = [
    dict(id="00_Resum", tab="00_Resum", titol="Resum executiu"),
    dict(id="A1_3_total", tab="A1_3_total", titol="A.1 — Despeses per concepte (total)"),
    dict(id="A1_1_sense", tab="A1_1_sense_afectacio", titol="A.1.1 — Despeses sense afectació"),
    dict(id="A1_2_amb", tab="A1_2_amb_afectacio", titol="A.1.2 — Despeses amb afectació"),
    dict(id="A2_0", tab="A2_0_resum_programa", titol="A.2 — Despeses per programa funcional"),
    dict(id="A3_0", tab="A3_0_resum_centre_capitol", titol="A.3 — Despeses per centre"),
    dict(id="B1_3", tab="B1_3_total", titol="B.1 — Ingressos per concepte (total)"),
    dict(id="B2_0", tab="B2_0_resum_tercer_capitol", titol="B.2 — Ingressos per entitat finançadora"),
    dict(id="C1_1", tab="C1_1_var_despeses_capitol", titol="C.1 — Variació de despeses per capítol"),
    dict(id="C2_1", tab="C2_1_var_ingressos_capitol", titol="C.2 — Variació d'ingressos per capítol"),
]
AMPLES = {"A3_0_resum_centre_capitol", "B2_0_resum_tercer_capitol"}
APAISADES = {"A3_0_resum_centre_capitol", "B2_0_resum_tercer_capitol"}


def _num(n, dec=2):
    try:
        x = float(n)
    except (TypeError, ValueError):
        return str(n)
    return f"{x:,.{dec}f}".replace(",", "§").replace(".", ",").replace("§", ".")


def fmt_pct(n):
    try:
        x = float(n)
    except (TypeError, ValueError):
        return str(n)
    return ("+" if x > 0 else "") + _num(x * 100, 2) + " %"


COLS_ETIQUETA = {"tipus_fila", "capitol", "article", "concepte", "fila", "programa_desc"}


def _es_pct(h):
    return str(h).startswith("pct") or "%" in str(h) or h == "variacio_pct"


def _es_num(v):
    # `numbers.Number` en el cuaderno; `numbers` no está en el ecosistema autorizado, y numpy sí.
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, (bool, np.bool_))


def _fmt_cel(h, v, dec):
    if v is None or v == "" or (isinstance(v, float) and pd.isna(v)):
        return ""
    if _es_num(v) and _es_pct(h):
        return fmt_pct(v)
    if _es_num(v) and h not in COLS_ETIQUETA:
        return _num(v, dec)
    return str(v)


LABELS = {
    "capitol": "Capítol", "article": "Article", "concepte": "Concepte",
    "import": "Import", "pct_cap": "Percentatge sobre capítol", "pct_total": "Percentatge sobre total",
    "fila": "Descripció", "Total": "Total", "% Total": "Percentatge sobre total",
    "anterior": "Exercici anterior", "actual": "Exercici actual",
    "variacio_eur": "Variació euros", "variacio_pct": "Variació percentatge",
    "programa_desc": "Programa funcional", "sense_afectacio": "Sense afectació",
    "amb_afectacio": "Amb afectació", "total": "Total",
}
_RE_CAP = re.compile(r"^\s*(\d+)\s*-\s*(.+)$")


def _label(h):
    if h in LABELS:
        return LABELS[h]
    s = str(h)
    return s[:1].upper() + s[1:] if s else s


_RE_VAL = re.compile(r"\{\{VAL:([A-Za-z0-9_]+)\}\}")


def calcular_valors(T):
    v = {}
    r = T["00_Resum"].set_index("concepte")
    v["total_despeses"] = _num(r.loc["Total despeses", "actual"], 0)
    v["total_ingressos"] = _num(r.loc["Total ingressos", "actual"], 0)
    v["var_despeses_pct"] = fmt_pct(r.loc["Total despeses", "variacio_pct"])
    v["var_despeses_eur"] = _num(r.loc["Total despeses", "variacio_eur"], 0)
    v["despeses_sense_afectacio"] = _num(r.loc["Despeses sense afectació", "actual"], 0)
    v["despeses_amb_afectacio"] = _num(r.loc["Despeses amb afectació", "actual"], 0)
    return v


def substitueix_valors(text, valors):
    """`{{VAL:clave}}` por su valor; si la clave no existe, deja `‹?clave›` a la vista."""
    if not valors:
        return text
    return _RE_VAL.sub(lambda m: valors.get(m.group(1), "‹?" + m.group(1) + "›"), text)


def llig_prosa(lineas):
    """La prosa de cada sección: lo que hay entre su título y su `{{TAULA:id:INICI}}`.

    Es la misma regla que el cuaderno aplicaba al Google Doc, sobre Markdown: un título (`#`)
    vacía lo acumulado, igual que allí lo vaciaba un párrafo con estilo de título.
    """
    prosa, buff = {}, []
    for linea in lineas:
        text = str(linea).strip()
        if not text:
            continue
        mi = re.match(r"\{\{TAULA:([A-Za-z0-9_]+):INICI\}\}", text)
        mf = re.match(r"\{\{TAULA:([A-Za-z0-9_]+):FI\}\}", text)
        if mi:
            prosa[mi.group(1)] = "\n".join(buff).strip()
            buff = []
            continue
        if mf or text.startswith("#"):
            buff = []
            continue
        buff.append(text)
    return prosa


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def _cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _camp(paragraph, instr):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " " + instr + " "
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


def _keep_with_next(par):
    pf = par.paragraph_format
    pf.keep_with_next = True
    pf.widow_control = True
    return par


def _ample_util(sec):
    return int(sec.page_width - sec.left_margin - sec.right_margin)


def _amplada_fixa(taula, ample_emu):
    taula.autofit = False
    n = len(taula.columns)
    if n > 1:
        primera = int(ample_emu * 0.22)
        resta = int((ample_emu - primera) / (n - 1))
        amples = [primera] + [resta] * (n - 1)
    else:
        amples = [ample_emu]
    for col, w in zip(taula.columns, amples):
        col.width = Emu(w)
        for cel in col.cells:
            cel.width = Emu(w)


def _afegir_taula(doc, df, cfg, ample_emu=None, abreviar_caps=False, mida_pt=None):
    vis = [c for c in df.columns if c != "tipus_fila"]
    taula = doc.add_table(rows=1, cols=len(vis))
    taula.style = "Table Grid"
    abrev = {}
    hdr = taula.rows[0]
    for j, c in enumerate(vis):
        etiqueta = _label(c)
        if abreviar_caps:
            m = _RE_CAP.match(str(c))
            if m:
                etiqueta = "Cap. " + m.group(1)
                abrev[etiqueta] = str(c)
        cell = hdr.cells[j]
        cell.text = etiqueta
        _shade(cell, cfg["header_fill"])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if mida_pt:
                    run.font.size = Pt(mida_pt)
    _repeat_header(hdr)
    _cant_split(hdr)
    for _, fila in df.iterrows():
        tipus = fila["tipus_fila"] if "tipus_fila" in df.columns else ""
        cels = taula.add_row().cells
        _cant_split(taula.rows[-1])
        for j, c in enumerate(vis):
            cel = cels[j]
            cel.text = _fmt_cel(c, fila[c], cfg["decimals"])
            p = cel.paragraphs[0]
            if c not in COLS_ETIQUETA:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                if tipus in ("total", "subtotal_cap"):
                    run.font.bold = True
                if tipus == "subtotal_art":
                    run.font.italic = True
                if mida_pt:
                    run.font.size = Pt(mida_pt)
    if ample_emu:
        _amplada_fixa(taula, ample_emu)
    return taula, abrev


def _afegir_llegenda(doc, abrev, mida_pt):
    if not abrev:
        return
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Abreviatures dels capítols (Cap. N = capítol N):")
    r.bold = True
    r.font.size = Pt(mida_pt)
    p.paragraph_format.space_after = Pt(0)
    for v in abrev.values():
        pv = doc.add_paragraph(v)
        pf = pv.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        for rr in pv.runs:
            rr.font.size = Pt(mida_pt)


def _config_seccio(sec, cfg, landscape):
    sec.top_margin = sec.bottom_margin = Cm(cfg["marge_cm"])
    sec.left_margin = sec.right_margin = Cm(cfg["marge_cm"])
    curt, llarg = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = llarg, curt
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = curt, llarg


def _toc(doc):
    p0 = doc.add_paragraph()
    r0 = p0.add_run("Índex")
    r0.bold = True
    r0.font.size = Pt(14)
    p = doc.add_paragraph()
    r = p.add_run()._r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r.append(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = r'TOC \o "1-2" \h \z \u'
    r.append(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r.append(sep)
    t = OxmlElement("w:t")
    t.text = "Actualitza l'índex: F9 o botó dret > Actualitza camp."
    r.append(t)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r.append(e)
    doc.add_page_break()


def _update_fields_on_open(doc):
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    doc.settings.element.append(el)


def escriu_docx_amb_prosa(T, prosa, path, cfg, valors=None):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = cfg["font"]
    normal.font.size = Pt(cfg["font_size"])
    _config_seccio(doc.sections[0], cfg, landscape=False)
    if cfg["encap_text"]:
        hp = doc.sections[0].header.paragraphs[0]
        hp.text = cfg["encap_text"]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Pàgina ")
    _camp(fp, "PAGE")
    fp.add_run(" de ")
    _camp(fp, "NUMPAGES")
    doc.add_heading(cfg["titol"], level=0)
    _toc(doc)
    n_prosa = 0
    for s in SECTIONS:
        if s["tab"] not in T:
            continue
        es_ampla = s["tab"] in AMPLES
        apaisar = s["tab"] in APAISADES
        if apaisar:
            _config_seccio(doc.add_section(WD_SECTION.NEW_PAGE), cfg, landscape=True)
        _keep_with_next(doc.add_heading(s["titol"], level=2))
        text = substitueix_valors(prosa.get(s["id"], "").strip(), valors)
        if text:
            n_prosa += 1
            for para in text.split("\n"):
                if para.strip():
                    _keep_with_next(doc.add_paragraph(para))
        else:
            pph = doc.add_paragraph()
            run = pph.add_run("[Text pendent d'inserir]")
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            _keep_with_next(pph)
        _keep_with_next(doc.add_paragraph())
        mida = cfg["mida_taula_ampla"] if es_ampla else None
        taula, abrev = _afegir_taula(doc, T[s["tab"]], cfg,
                                     ample_emu=_ample_util(doc.sections[-1]) if es_ampla else None,
                                     abreviar_caps=es_ampla, mida_pt=mida)
        if abrev:
            _afegir_llegenda(doc, abrev, mida or cfg["font_size"])
        if apaisar:
            _config_seccio(doc.add_section(WD_SECTION.NEW_PAGE), cfg, landscape=False)
        else:
            doc.add_paragraph()
    _update_fields_on_open(doc)
    doc.save(path)
    return n_prosa


def _lineas(ruta):
    # El auditor no deja abrir ficheros a mano: un fichero de texto se lee con una librería. El
    # separador es un carácter de control que el texto no lleva, para que cada línea sea una fila
    # entera; con el tabulador, una línea de Markdown con un tabulador se partiría en columnas.
    tabla = pd.read_csv(ruta, header=None, names=["linea"], sep="\x1f", quoting=3, dtype=str,
                        keep_default_na=False, encoding="utf-8-sig", skip_blank_lines=True)
    return tabla["linea"].tolist()


# =============================================================================
#  EJECUCIÓN
# =============================================================================
rutas = options["ficheros"]
cfg = dict(
    titol=options.get("titulo") or "Pressupost UJI — proposta",
    encap_text=options.get("encabezado") or "Universitat Jaume I — Pressupost",
    font="Calibri", font_size=10, marge_cm=1.5, decimals=2,
    mida_taula_ampla=8, header_fill="0B5394",
)

df, ter_tot = fase1_normalitza(rutas["presupuesto"], rutas["clasificacion"])
df = terceros(df, ter_tot)
salida_csv = output_dir + "/pressupost_net.csv"
df.to_csv(salida_csv, index=False, encoding="utf-8-sig")

T = genera_taules(salida_csv)
prosa = llig_prosa(_lineas(rutas["prosa"])) if rutas.get("prosa") else {}
valors = calcular_valors(T)
n_prosa = escriu_docx_amb_prosa(T, prosa, output_dir + "/pressupost_uji.docx", cfg, valors)

act = df[df["es_actual"]]
result = {
    "metrics": [
        {"name": "filas", "value": len(df)},
        {"name": "tablas", "value": len(T)},
        {"name": "secciones_con_prosa", "value": n_prosa},
        {"name": "total_gastos", "value": round(float(act[act["es_despesa"]]["importe"].sum()), 2)},
        {"name": "total_ingresos", "value": round(float(act[act["es_ingres"]]["importe"].sum()), 2)},
    ]
}
